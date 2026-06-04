from pptx import Presentation
from pptx.util import Inches, Pt

def crear_presentacion():
    prs = Presentation()
    
    # Diapositiva 1: Portada
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.shapes.placeholders[1]
    title.text = "Redes Neuronales Artificiales y Clasificación de Imágenes mediante CNN"
    subtitle.text = "Seminario Práctico\n\nAutor: Miguel Ángel Banteurt Blanco\nAsignatura: Inteligencia Artificial"
    
    # Diapositiva 2: Introducción
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.shapes.placeholders[1]
    title.text = "Introducción"
    tf = content.text_frame
    tf.text = "Las Redes Neuronales Artificiales (RNA) son una de las tecnologías más importantes en IA moderna"
    p = tf.add_paragraph()
    p.text = "Inspiradas en el cerebro humano"
    p = tf.add_paragraph()
    p.text = "Aplicaciones: reconocimiento de imágenes, NLP, sistemas de recomendación, conducción autónoma, diagnóstico médico"
    
    # Diapositiva 3: Conceptos Fundamentales
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.shapes.placeholders[1]
    title.text = "Conceptos Fundamentales"
    tf = content.text_frame
    tf.text = "Neurona Artificial: y = f(Σ(wᵢxᵢ) + b)"
    p = tf.add_paragraph()
    p.text = "Funciones de Activación:"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "- ReLU: f(x) = max(0,x)"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "- Sigmoid: valores entre 0 y 1"
    p.level = 2
    p = tf.add_paragraph()
    p.text = "- Softmax: probabilidades para clasificación multiclase"
    p.level = 2
    
    # Diapositiva 4: Arquitecturas
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.shapes.placeholders[1]
    title.text = "Arquitecturas de Redes Neuronales"
    tf = content.text_frame
    tf.text = "CNN (Convolutional Neural Networks)"
    p = tf.add_paragraph()
    p.text = "Especializadas en imágenes"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "RNN (Recurrent Neural Networks)"
    p = tf.add_paragraph()
    p.text = "Datos secuenciales con memoria temporal"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "LSTM (Long Short-Term Memory)"
    p = tf.add_paragraph()
    p.text = "Secuencias largas con puertas de memoria"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Transformers"
    p = tf.add_paragraph()
    p.text = "Mecanismos de atención, procesamiento paralelo"
    p.level = 1
    
    # Diapositiva 5: Dataset MNIST
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.shapes.placeholders[1]
    title.text = "Dataset MNIST"
    tf = content.text_frame
    tf.text = "70,000 imágenes en escala de grises"
    p = tf.add_paragraph()
    p.text = "Resolución: 28 × 28 píxeles"
    p = tf.add_paragraph()
    p.text = "Dígitos del 0 al 9 (10 clases)"
    p = tf.add_paragraph()
    p.text = "Estándar para introducir Deep Learning"
    
    # Diapositiva 6: Resultados Esperados
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.shapes.placeholders[1]
    title.text = "Resultados Esperados"
    tf = content.text_frame
    tf.text = "Después de 5 épocas de entrenamiento:"
    p = tf.add_paragraph()
    p.text = "Precisión de entrenamiento: 98% – 99%"
    p = tf.add_paragraph()
    p.text = "Precisión de prueba: 97% – 99%"
    p = tf.add_paragraph()
    p.text = "Estos resultados evidencian la efectividad de las CNN"
    
    # Guardar presentación
    prs.save('Presentacion.pptx')
    print("✅ Presentacion.pptx generada exitosamente")

if __name__ == "__main__":
    crear_presentacion()
